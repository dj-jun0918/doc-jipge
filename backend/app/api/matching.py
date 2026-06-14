import csv
import io
import logging
import uuid
from collections import Counter
from datetime import datetime
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.database import get_db
from app.matcher.matcher import (
    compute_aggregate_score,
    compute_field_sensitivities,
    counterfactual_for_field,
    derive_eligibility_bucket,
    match_announcement,
)
from app.models.announcement import Announcement
from app.models.company import Company
from app.models.eligibility import EligibilityResult
from app.models.match_result import MatchResult
from app.schemas.eligibility import EligibilityField, ParsedCondition
from app.schemas.matching import (
    BucketCounts,
    CompanyMatchListResponse,
    CompanyMatchSummary,
    CounterfactualItem,
    CounterfactualRequest,
    CounterfactualResponse,
    MatchResultDetailResponse,
    MatchResultItem,
    MatchResultStats,
    SimulateOverrides,
    SimulateRequest,
    SimulateResponse,
    TriggerResponse,
)

logger = logging.getLogger(__name__)

router = APIRouter()

# 버킷 정렬 우선순위 — 신청가능 → 조건확인 → 자격미달
_BUCKET_RANK = {"신청가능": 0, "조건확인": 1, "자격미달": 2}


def _parse_uuid(value: str, what: str) -> uuid.UUID:
    try:
        return uuid.UUID(value)
    except ValueError:
        raise HTTPException(status_code=400, detail=f"{what} UUID 형식 오류: {value}")


@router.get("/{company_id}", response_model=CompanyMatchListResponse)
def get_matching_results(
    company_id: str,
    limit: int = Query(10, ge=1, le=50),
    db: Session = Depends(get_db),
):
    """회사 전체 매칭 결과 — 가중 합산 총점(연속 score) 기준 정렬 (TOP N)."""
    company_uuid = _parse_uuid(company_id, "company_id")

    if not db.get(Company, company_uuid):
        raise HTTPException(status_code=404, detail="회사를 찾을 수 없습니다")

    rows = db.execute(
        select(
            MatchResult.announcement_id,
            MatchResult.field_name,
            MatchResult.status,
            MatchResult.score,
        )
        .where(MatchResult.company_id == company_uuid)
    ).all()

    by_ann: dict[uuid.UUID, list[tuple[str, str, float | None]]] = {}
    for ann_id, field_name, status, score in rows:
        by_ann.setdefault(ann_id, []).append((field_name, status, score))

    summaries: list[tuple[uuid.UUID, float, int, int, str]] = []
    counts = {"신청가능": 0, "조건확인": 0, "자격미달": 0}
    for ann_id, fields in by_ann.items():
        agg_score = compute_aggregate_score(fields)
        fulfilled = sum(1 for _, status, _ in fields if status == "충족")
        total = len(fields)
        bucket = derive_eligibility_bucket(fields)
        counts[bucket] += 1
        summaries.append((ann_id, agg_score, fulfilled, total, bucket))

    # 버킷 우선(신청가능→조건확인→자격미달), 버킷 내 점수 내림차순 — "확실한 탈락"이
    # "확인하면 될 수도 있는 공고"보다 위에 오지 않도록. 자격미달은 점수순이라 "거의 됨"이 위로.
    summaries.sort(key=lambda x: (_BUCKET_RANK[x[4]], -x[1]))
    top = summaries[:limit]

    ann_ids = [s[0] for s in top]
    titles: dict[uuid.UUID, str] = {}
    if ann_ids:
        for ann in db.scalars(select(Announcement).where(Announcement.id.in_(ann_ids))):
            titles[ann.id] = ann.title

    items = [
        CompanyMatchSummary(
            announcement_id=ann_id,
            title=titles.get(ann_id, ""),
            match_score=round(score, 3),
            fulfilled_count=fulfilled,
            total_fields=total,
            bucket=bucket,
        )
        for ann_id, score, fulfilled, total, bucket in top
    ]

    return CompanyMatchListResponse(
        company_id=company_uuid,
        items=items,
        total=len(summaries),
        bucket_counts=BucketCounts(**counts),
    )


@router.get("/{company_id}/export")
def export_matching_results(
    company_id: str,
    format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
    db: Session = Depends(get_db),
):
    """매칭 결과를 Excel/CSV 파일로 다운로드."""
    company_uuid = _parse_uuid(company_id, "company_id")
    company = db.get(Company, company_uuid)
    if not company:
        raise HTTPException(status_code=404, detail="기업을 찾을 수 없습니다")

    stmt = (
        select(MatchResult, Announcement)
        .join(Announcement, MatchResult.announcement_id == Announcement.id)
        .where(MatchResult.company_id == company_uuid)
        .order_by(MatchResult.created_at.desc())
    )
    rows = db.execute(stmt).all()

    if format == "xlsx":
        return _export_xlsx(company, rows)
    return _export_csv(company, rows)


@router.get(
    "/{company_id}/{announcement_id}",
    response_model=MatchResultDetailResponse,
)
def get_matching_detail(
    company_id: str,
    announcement_id: str,
    db: Session = Depends(get_db),
):
    """특정 공고 vs 회사 필드별 매칭 결과."""
    company_uuid = _parse_uuid(company_id, "company_id")
    ann_uuid = _parse_uuid(announcement_id, "announcement_id")

    rows = db.scalars(
        select(MatchResult)
        .where(
            MatchResult.company_id == company_uuid,
            MatchResult.announcement_id == ann_uuid,
        )
        .order_by(MatchResult.created_at)
    ).all()

    if not rows:
        return MatchResultDetailResponse(
            company_id=company_uuid,
            announcement_id=ann_uuid,
            items=[],
            stats=MatchResultStats(),
            matched_at=None,
        )

    items = [
        MatchResultItem(
            field_name=r.field_name,
            status=r.status,
            score=r.score,
            distance=r.distance,
            constraint_type=r.constraint_type,
            company_value=r.company_value,
            requirement_value=r.requirement_value,
            evidence=r.evidence,
            processing_path=r.processing_path,
        )
        for r in rows
    ]
    counter = Counter(r.status for r in rows)
    stats = MatchResultStats(
        충족=counter.get("충족", 0),
        미충족=counter.get("미충족", 0),
        확인필요=counter.get("확인필요", 0),
        해당없음=counter.get("해당없음", 0),
    )

    field_tuples = [(r.field_name, r.status, r.score) for r in rows]
    return MatchResultDetailResponse(
        company_id=company_uuid,
        announcement_id=ann_uuid,
        items=items,
        stats=stats,
        match_score=round(compute_aggregate_score(field_tuples), 3),
        bucket=derive_eligibility_bucket(field_tuples),
        matched_at=rows[-1].created_at,
    )


def _eligibility_fields(elig_rows: list[EligibilityResult]) -> list[EligibilityField]:
    """EligibilityResult 행들을 매칭 입력 EligibilityField로 변환."""
    return [
        EligibilityField(
            field_name=er.field_name,
            condition=ParsedCondition(
                value=(er.condition_parsed or {}).get("value"),
                operator=(er.condition_parsed or {}).get("operator"),
                raw_text=er.condition_value,
            ),
            evidence=er.evidence or "",
            evidence_source=er.evidence_source or "",
            processing_path=er.processing_path,
        )
        for er in elig_rows
    ]


def _apply_overrides(company: Company, overrides: SimulateOverrides) -> Company:
    """Company의 detached copy 생성 후 overrides 적용 (None 아닌 필드만). session add X."""
    return Company(
        id=company.id,
        name=company.name,
        founded_date=overrides.founded_date or company.founded_date,
        revenue=overrides.revenue if overrides.revenue is not None else company.revenue,
        region=overrides.region if overrides.region is not None else company.region,
        industry=overrides.industry if overrides.industry is not None else company.industry,
        employee_count=(
            overrides.employee_count
            if overrides.employee_count is not None
            else company.employee_count
        ),
        ceo_birth_date=overrides.ceo_birth_date or company.ceo_birth_date,
        certifications=(
            overrides.certifications
            if overrides.certifications is not None
            else company.certifications
        ),
    )


@router.post("/{company_id}/simulate", response_model=SimulateResponse)
def simulate_matching(
    company_id: str,
    req: SimulateRequest,
    db: Session = Depends(get_db),
):
    """What-if 시뮬레이션 — 회사 프로필 임시 변경값으로 매칭 재계산 (DB 저장 X)."""
    company_uuid = _parse_uuid(company_id, "company_id")
    company = db.get(Company, company_uuid)
    if not company:
        raise HTTPException(status_code=404, detail="회사를 찾을 수 없습니다")

    elig_rows = db.scalars(
        select(EligibilityResult).where(
            EligibilityResult.announcement_id == req.announcement_id
        )
    ).all()

    if not elig_rows:
        return SimulateResponse(
            company_id=company_uuid,
            announcement_id=req.announcement_id,
            items=[],
            stats=MatchResultStats(),
            matched_at=None,
        )

    fields = _eligibility_fields(elig_rows)

    sim_company = _apply_overrides(company, req.overrides)
    results = match_announcement(sim_company, fields, req.announcement_id)

    items = [
        MatchResultItem(
            field_name=r.field_name,
            status=r.status,
            score=r.score,
            distance=r.distance,
            constraint_type=r.constraint_type,
            company_value=r.company_value,
            requirement_value=r.requirement_value,
            evidence=r.evidence,
            processing_path=r.processing_path,
        )
        for r in results
    ]
    counter = Counter(r.status for r in results)
    stats = MatchResultStats(
        충족=counter.get("충족", 0),
        미충족=counter.get("미충족", 0),
        확인필요=counter.get("확인필요", 0),
        해당없음=counter.get("해당없음", 0),
    )

    field_tuples = [(r.field_name, r.status, r.score) for r in results]
    return SimulateResponse(
        company_id=company_uuid,
        announcement_id=req.announcement_id,
        items=items,
        stats=stats,
        match_score=round(compute_aggregate_score(field_tuples), 3),
        bucket=derive_eligibility_bucket(field_tuples),
        matched_at=None,  # 시뮬레이션은 저장 X
    )


@router.post("/{company_id}/counterfactual", response_model=CounterfactualResponse)
def counterfactual_matching(
    company_id: str,
    req: CounterfactualRequest,
    db: Session = Depends(get_db),
):
    """반사실 분석 — 미충족 공고를 충족시키는 최소 프로필 변경 제안 (DB 저장 X).

    제안은 변경 가능 조건을 먼저, 그 안에서 영향 큰(총점 상승폭 큰) 조건 순으로 정렬한다.
    """
    company_uuid = _parse_uuid(company_id, "company_id")
    company = db.get(Company, company_uuid)
    if not company:
        raise HTTPException(status_code=404, detail="회사를 찾을 수 없습니다")

    elig_rows = db.scalars(
        select(EligibilityResult).where(
            EligibilityResult.announcement_id == req.announcement_id
        )
    ).all()
    if not elig_rows:
        return CounterfactualResponse(
            company_id=company_uuid,
            announcement_id=req.announcement_id,
            unmet=[],
            achievable=True,
            note="자격요건이 없는 공고입니다",
        )

    fields = _eligibility_fields(elig_rows)
    results = match_announcement(company, fields, req.announcement_id)

    unmet: list[CounterfactualItem] = []
    override_kwargs: dict = {}
    has_unchangeable = False

    for field, r in zip(fields, results):
        if r.status != "미충족":
            continue
        cf = counterfactual_for_field(field, company)
        if cf is None:
            continue
        unmet.append(CounterfactualItem(
            field_name=field.field_name,
            current_value=r.company_value,
            requirement=r.requirement_value,
            suggested_value=cf["suggested_value"],
            explanation=cf["explanation"],
            changeable=cf["changeable"],
        ))
        if cf["changeable"] and cf["override_attr"]:
            override_kwargs[cf["override_attr"]] = cf["override_value"]
        else:
            has_unchangeable = True

    # 영향 큰 조건부터 — 변경 가능 조건을 먼저, 그 안에서 sensitivity(총점 상승폭) 내림차순
    sens = compute_field_sensitivities(
        [(r.field_name, r.status, r.score) for r in results]
    )
    unmet.sort(key=lambda it: (it.changeable, sens.get(it.field_name, 0.0)), reverse=True)

    # 달성 가능 여부 — changeable 변경 모두 적용 후 미충족이 없어야 함
    if not unmet:
        achievable = True
    elif has_unchangeable:
        achievable = False  # 변경 불가 조건이 미충족으로 남음
    else:
        sim_company = _apply_overrides(company, SimulateOverrides(**override_kwargs))
        verified = match_announcement(sim_company, fields, req.announcement_id)
        achievable = all(v.status != "미충족" for v in verified)

    note = (
        "일부 조건(업력/나이/업종 제외 등)은 프로필 변경으로 충족 불가"
        if has_unchangeable
        else None
    )

    return CounterfactualResponse(
        company_id=company_uuid,
        announcement_id=req.announcement_id,
        unmet=unmet,
        achievable=achievable,
        note=note,
    )


@router.post("/{company_id}/run", response_model=TriggerResponse)
def trigger_matching(company_id: str, db: Session = Depends(get_db)):
    """회사 vs DB 내 자격요건 추출된 공고 전체 매칭 — Celery 비동기."""
    from app.worker.tasks import match_company_announcements

    company_uuid = _parse_uuid(company_id, "company_id")

    if not db.get(Company, company_uuid):
        raise HTTPException(status_code=404, detail="회사를 찾을 수 없습니다")

    task = match_company_announcements.delay(str(company_uuid))
    return TriggerResponse(task_id=task.id, message="매칭 작업이 시작되었습니다.")


# ── Export helpers ────────────────────────────────────────────────────────────

_EXPORT_HEADERS = [
    "공고명", "필드", "조건", "회사 값", "판정",
    "점수", "거리", "제약", "근거", "처리 경로",
]
_COL_WIDTHS = [50, 20, 30, 30, 10, 10, 10, 10, 50, 15]


def _row_values(match: MatchResult, ann: Announcement) -> list:
    return [
        ann.title,
        match.field_name,
        match.requirement_value or "",
        match.company_value or "",
        match.status,
        round(match.score, 3) if match.score is not None else "",
        round(match.distance, 3) if match.distance is not None else "",
        match.constraint_type or "",
        (match.evidence or {}).get("text", ""),  # evidence JSONB에서 표시용 text만
        match.processing_path,
    ]


def _content_disposition(filename: str) -> str:
    """RFC 5987 방식으로 Content-Disposition 헤더 값 반환 (한글 파일명 안전)."""
    encoded = quote(filename)
    return f"attachment; filename*=UTF-8''{encoded}"


def _export_xlsx(company: Company, rows: list) -> StreamingResponse:
    """openpyxl로 Excel 생성 + StreamingResponse."""
    wb = Workbook()
    ws = wb.active
    ws.title = "매칭 결과"

    # 헤더 스타일 (파란 배경 + 흰 글씨 + 가운데 정렬)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2563EB")
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, text in enumerate(_EXPORT_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # 데이터 행
    for match, ann in rows:
        ws.append(_row_values(match, ann))

    # 컬럼 너비 지정
    for col_idx, width in enumerate(_COL_WIDTHS, start=1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"matching_{company.name}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(filename)},
    )


def _export_csv(company: Company, rows: list) -> StreamingResponse:
    """csv 모듈로 CSV 생성 + StreamingResponse (UTF-8 BOM 포함)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_EXPORT_HEADERS)
    for match, ann in rows:
        writer.writerow(_row_values(match, ann))

    # \ufeff = UTF-8 BOM — Excel에서 CSV 열 때 한글 깨짐 방지
    content = "\ufeff" + buffer.getvalue()
    filename = f"matching_{company.name}_{datetime.now().strftime('%Y%m%d')}.csv"
    return StreamingResponse(
        io.BytesIO(content.encode("utf-8")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": _content_disposition(filename)},
    )
